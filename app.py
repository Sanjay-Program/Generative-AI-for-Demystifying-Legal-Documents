# legal_ai_dashboard_final.py
# Single-file FastAPI app that loads laws from sample_laws.json
# Run: python legal_ai_dashboard_final.py

import os
import io
import re
import asyncio
import json
from typing import List, Optional
from html import unescape
from contextlib import asynccontextmanager

from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse
from pydantic import BaseModel, field_validator
import uvicorn

# AI
import google.generativeai as genai

# File handling
import docx
import fitz  # PyMuPDF

# PDF export
from fpdf import FPDF

# Database (SQLAlchemy)
from sqlalchemy import create_engine, Column, Integer, String, Text, DateTime, func
from sqlalchemy.orm import sessionmaker, declarative_base, Session

# --- 0. CONFIG ---
load_dotenv()
API_KEY = os.getenv("GOOGLE_API_KEY", "")
if not API_KEY:
    print("⚠️ GOOGLE_API_KEY not found in environment. Make sure to set it in .env")

try:
    genai.configure(api_key=API_KEY)
    model = genai.GenerativeModel('gemini-1.5-flash')
    print("✅ Successfully configured Google AI.")
except Exception as e:
    print(f"🔴 CRITICAL ERROR: Failed to configure Google AI. Check API key. Details: {e}")

# --- 1. DATABASE SETUP (SQLite + SQLAlchemy) ---
DB_PATH = os.getenv("LAW_DB_PATH", "laws.db")
engine = create_engine(f"sqlite:///{DB_PATH}", connect_args={"check_same_thread": False})
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
Base = declarative_base()


class Law(Base):
    __tablename__ = "laws"
    id = Column(Integer, primary_key=True, index=True)
    title = Column(String(400), nullable=False)
    jurisdiction = Column(String(120), default="India")
    tags = Column(String(400), default="")
    text = Column(Text, nullable=False)
    created_at = Column(DateTime(timezone=True), server_default=func.now())


class SuggestionRecord(Base):
    __tablename__ = "suggestions"
    id = Column(Integer, primary_key=True, index=True)
    source_doc_preview = Column(String(800), nullable=True)
    suggested_text = Column(Text, nullable=False)
    reason = Column(Text, nullable=True)
    created_at = Column(DateTime(timezone=True), server_default=func.now())


Base.metadata.create_all(bind=engine)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

# --- 2. UTILITIES & HELPERS ---
def sanitize_for_pdf(text: str) -> str:
    return text.encode("latin-1", "replace").decode("latin-1")


def extract_text_from_file(file: UploadFile) -> str:
    filename = file.filename or ""
    try:
        if filename.lower().endswith(".pdf"):
            data = file.file.read()
            with fitz.open(stream=data, filetype="pdf") as doc:
                return "\n".join([page.get_text() for page in doc])
        elif filename.lower().endswith(".docx"):
            data = file.file.read()
            docx_doc = docx.Document(io.BytesIO(data))
            return "\n".join([p.text for p in docx_doc.paragraphs])
        elif filename.lower().endswith(".txt"):
            return file.file.read().decode("utf-8", errors="ignore")
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type. Use .pdf, .docx or .txt")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to process file: {e}")


async def get_ai_response(prompt: str) -> str:
    try:
        if not API_KEY:
            return "AI API key not configured."
        out = await model.generate_content_async(prompt)
        return out.text if hasattr(out, "text") else str(out)
    except Exception as e:
        print("AI Error:", e)
        return f"AI error: {e}"

# --- MODIFIED: This function now reads from sample_laws.json ---
def seed_sample_laws(db: Session):
    """Seed DB from sample_laws.json if the DB is empty."""
    if db.query(Law).count() > 0:
        return

    try:
        with open("sample_laws.json", "r", encoding="utf-8") as f:
            sample_laws = json.load(f)
        
        for s in sample_laws:
            db.add(Law(
                title=s.get("title", "No Title"),
                jurisdiction=s.get("jurisdiction", "India"),
                tags=s.get("tags", ""),
                text=s.get("text", "")
            ))
        db.commit()
        print(f"✅ Seeded {len(sample_laws)} sample laws into DB from sample_laws.json.")
    except FileNotFoundError:
        print("⚠️ 'sample_laws.json' not found. No laws were seeded.")
    except json.JSONDecodeError:
        print("🔴 Error decoding 'sample_laws.json'. Please check if it is a valid JSON file.")
    except Exception as e:
        print(f"🔴 An unexpected error occurred while seeding laws: {e}")


def simple_law_search(db: Session, q: Optional[str] = None, jurisdiction: Optional[str] = None, limit: int = 10):
    qexpr = f"%{q.strip()}%" if q else "%"
    query = db.query(Law)
    if jurisdiction:
        query = query.filter(Law.jurisdiction.ilike(f"%{jurisdiction}%"))
    if q:
        query = query.filter((Law.title.ilike(qexpr)) | (Law.tags.ilike(qexpr)) | (Law.text.ilike(qexpr)))
    return query.order_by(Law.created_at.desc()).limit(limit).all()

# --- 3. Pydantic Schemas ---
class AskRequest(BaseModel):
    document_text: str
    question: str
    language: str

class NegotiateRequest(BaseModel):
    history: list
    user_message: str

class CompareRequest(BaseModel):
    clause_a: str
    clause_b: str
    language: str

class ReportRequest(BaseModel):
    key_facts: str
    risk_analysis: str
    filled_document: str
    @field_validator('*', mode='before')
    def sanitize_fields(cls, v):
        if isinstance(v, str): return sanitize_for_pdf(v)
        return v

class LawSearchRequest(BaseModel):
    document_text: Optional[str] = ""
    q: Optional[str] = ""
    language: Optional[str] = "en"
    jurisdiction: Optional[str] = "Chennai"

# --- 4. PDF Helper ---
class PDF(FPDF):
    def header(self):
        self.set_font('Helvetica', 'B', 15)
        self.cell(0, 10, 'Legal AI Analysis Report', 0, 1, 'C')
        self.ln(5)
    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', 0, 0, 'C')
    def chapter_title(self, title):
        self.ln(5)
        self.set_font('Helvetica', 'B', 12)
        self.cell(0, 10, sanitize_for_pdf(title), 0, 1, 'L')
        self.ln(2)
    def write_html_content(self, html):
        clean_html = re.sub('<[^<]+?>', '', unescape(html))
        self.set_font('Helvetica', '', 10)
        for line in clean_html.split('\n'):
            if line.strip(): self.multi_cell(0, 5, sanitize_for_pdf(line.strip()))
        self.ln()

# --- 5. FRONTEND (HTML, CSS, JS) ---
HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title data-translate-key="title">Advanced Legal AI Dashboard</title>
    <link href="https://fonts.googleapis.com/css2?family=Poppins:wght@300;400;600;700&display=swap" rel="stylesheet">
    <style>{{ css_code }}</style>
</head>
<body>
    <canvas id="particle-canvas"></canvas>
    <div id="splash-screen">
        <div class="logo">⚖️</div>
        <h1 data-translate-key="splash_title">Generative AI for Legal Documents</h1>
        <p data-translate-key="splash_subtitle">Loading Advanced Analysis Suite...</p>
    </div>

    <main id="app-body" class="hidden">
        <header>
            <h1 data-translate-key="header_title">Advanced Legal AI Dashboard</h1>
            <p data-translate-key="header_subtitle">Empowering You to Make Informed Decisions with Confidence</p>
             <div class="settings-bar">
                 <select id="languageSelector">
                     <option value="en">English</option>
                     <option value="hi">हिन्दी (Hindi)</option>
                     <option value="ta">தமிழ் (Tamil)</option>
                     <option value="te">తెలుగు (Telugu)</option>
                     <option value="ml">മലയാളം (Malayalam)</option>
                     <option value="kn">ಕನ್ನಡ (Kannada)</option>
                     <option value="es">Español (Spanish)</option>
                     <option value="fr">Français (French)</option>
                 </select>
             </div>
        </header>

        <div class="container">
            <div class="document-input-card card-glow">
                <h2 data-translate-key="begin_analysis_title">1. Begin Your Analysis</h2>
                <div class="file-upload-area">
                    <input type="file" id="documentUpload" accept=".txt,.pdf,.docx" class="hidden-file-input">
                    <label for="documentUpload" class="file-upload-label" data-translate-key="upload_label">Click to Upload Document (.txt, .pdf, .docx)</label>
                    <span id="fileName" class="file-name-display"></span>
                </div>
                <div class="autofill-section">
                    <input type="text" id="autofillName" data-translate-key="name_placeholder" placeholder="Enter Your Full Name for Auto-Filling">
                    <button id="analyzeBtn" data-translate-key="analyze_button">🚀 Analyze & Build Dashboard</button>
                </div>
            </div>

            <div id="loader" class="hidden"></div>

            <div id="dashboard" class="hidden">
                 <div class="dashboard-header">
                    <h2 data-translate-key="analysis_results">Analysis Results</h2>
                    <button id="downloadReportBtn" data-translate-key="download_button">📄 Download Report</button>
                </div>
                <div class="dashboard-grid">
                    <div class="card card-glow key-facts">
                        <h3 class="card-header" data-translate-key="card_key_facts">📊 Key Facts & Figures</h3>
                        <div id="keyFactsOutput" class="card-content"></div>
                    </div>
                    <div class="card card-glow risk-analysis">
                        <h3 class="card-header" data-translate-key="card_risk_analysis">🚦 Risk Analysis</h3>
                        <div id="riskAnalysisOutput" class="card-content"></div>
                    </div>
                    <div class="card card-glow legal-lifespan">
                        <h3 class="card-header" data-translate-key="card_lifespan">🗓️ Legal Lifespan & Timeline</h3>
                        <div id="lifespanOutput" class="card-content"></div>
                    </div>
                     <div class="card card-glow clause-comparison">
                        <h3 class="card-header" data-translate-key="card_clause_comparison">🔍 Clause-by-Clause Analysis</h3>
                        <div class="comparison-inputs">
                            <textarea id="clauseA" data-translate-key="clause_a_placeholder" placeholder="Paste Original Clause Here..."></textarea>
                            <textarea id="clauseB" data-translate-key="clause_b_placeholder" placeholder="Paste Modified Clause Here..."></textarea>
                        </div>
                        <button id="compareBtn" data-translate-key="compare_button">Compare Clauses</button>
                        <div id="comparisonOutput" class="card-content result-box"></div>
                    </div>
                    <div class="card card-glow interactive-qa">
                        <h3 class="card-header" data-translate-key="card_qa">💬 Interactive Q&A (Voice Enabled)</h3>
                        <div id="qaOutput" class="card-content chat-box"></div>
                        <div class="qa-input">
                            <input type="text" id="qaInput" data-translate-key="qa_placeholder" placeholder="Ask a question...">
                            <button id="askBtn" title="Send Question">➤</button>
                            <button id="speakBtn" title="Ask with your voice">🎙️</button>
                        </div>
                    </div>
                    <div class="card card-glow negotiation-simulator">
                        <h3 class="card-header" data-translate-key="card_negotiation">🤝 AI Negotiation Simulator</h3>
                        <div id="negotiationOutput" class="card-content chat-box"></div>
                        <div class="qa-input">
                            <input type="text" id="negotiationInput" data-translate-key="negotiation_placeholder" placeholder="Type your negotiation point...">
                            <button id="negotiateBtn" title="Send Negotiation Point">➤</button>
                        </div>
                    </div>
                     <div class="card card-glow legal-aid">
                        <h3 class="card-header" data-translate-key="card_legal_aid">👨‍⚖️ Local Laws & Suggestions (Chennai)</h3>
                        <button id="legalAidBtn" data-translate-key="legal_aid_button">Find Relevant Laws & Suggestions</button>
                        <div id="legalAidOutput" class="card-content result-box"></div>
                    </div>
                    <div class="card card-glow full-span-card filled-document">
                        <h3 class="card-header" data-translate-key="card_autofilled">📝 Your Auto-Filled Document</h3>
                        <pre id="filledDocumentOutput" class="card-content code-box"></pre>
                    </div>
                </div>
            </div>
        </div>
    </main>
    <footer>
        <p data-translate-key="built_by">Built by Visoneers. For updates or changes, contact: <a href="mailto:sanjaymurugadoss02@gmail.com">sanjaymurugadoss02@gmail.com</a></p>
        <p data-translate-key="disclaimer">Disclaimer: This AI tool is for informational purposes only and is not a substitute for professional legal advice.</p>
    </footer>
    <script>{{ js_code }}</script>
</body>
</html>
"""

CSS_CODE = """
:root {
    --bg-color: #0F0F1A;
    --primary-accent: #8A2BE2; /* BlueViolet */
    --secondary-accent: #00BFFF; /* DeepSkyBlue */
    --glow-color: rgba(138, 43, 226, 0.6);
    --text-color: #EAEAEA;
    --text-muted: #A0A0A0;
    --font-family: 'Poppins', sans-serif;
    --card-bg: rgba(22, 22, 34, 0.6);
    --high-risk-bg: rgba(233, 69, 96, 0.25);
    --caution-bg: rgba(247, 183, 49, 0.25);
}
* { box-sizing: border-box; margin: 0; padding: 0; }
body {
    font-family: var(--font-family); background-color: var(--bg-color);
    color: var(--text-color); padding: 20px;
    line-height: 1.7; overflow-x: hidden;
}
#particle-canvas {
    position: fixed; top: 0; left: 0;
    z-index: -1; width: 100%; height: 100%;
}
#splash-screen {
    position: fixed; top: 0; left: 0; width: 100%; height: 100%;
    background: var(--bg-color); display: flex; flex-direction: column;
    justify-content: center; align-items: center; z-index: 1000;
    transition: opacity 0.5s ease-out;
}
#splash-screen .logo { font-size: 5rem; animation: pulse 2s infinite; }
.container { max-width: 1800px; margin: 0 auto; }
header { text-align: center; margin-bottom: 20px; animation: fadeInDown 1s ease-out; }
h1 {
    background: linear-gradient(90deg, var(--primary-accent), var(--secondary-accent));
    -webkit-background-clip: text;
    -webkit-text-fill-color: transparent;
    font-weight: 700; font-size: 2.8rem; margin-bottom: 10px;
}
.settings-bar { margin-bottom: 20px; }
#languageSelector {
    padding: 10px 15px; border-radius: 8px; border: 1px solid var(--primary-accent);
    background-color: var(--card-bg); color: var(--text-color); font-size: 1rem;
    cursor: pointer; transition: box-shadow 0.3s ease;
}
#languageSelector:focus { outline: none; box-shadow: 0 0 15px var(--glow-color); }
.card, .document-input-card {
    background-color: var(--card-bg); backdrop-filter: blur(10px); -webkit-backdrop-filter: blur(10px);
    border-radius: 20px; padding: 30px; margin-bottom: 25px;
    border: 1px solid rgba(138, 43, 226, 0.3);
    box-shadow: 0 8px 32px 0 rgba(0, 0, 0, 0.37);
    transition: transform 0.3s ease, border-color 0.3s ease;
    animation: card-fade-in 0.6s ease-out forwards;
}
.card-glow:hover { transform: translateY(-8px); border-color: rgba(138, 43, 226, 0.8); }
.file-upload-label {
    display: block; padding: 20px; border-radius: 10px;
    border: 2px dashed rgba(138, 43, 226, 0.5); background-color: rgba(0,0,0,0.2);
    color: var(--text-muted); text-align: center; cursor: pointer;
    transition: background-color 0.3s, border-color 0.3s;
}
.file-upload-label:hover { border-color: var(--secondary-accent); color: #fff; }
.file-name-display { display: block; text-align: center; margin-top: 10px; color: var(--secondary-accent); }
input[type="text"], textarea {
    width: 100%; padding: 15px; border-radius: 10px;
    border: 1px solid rgba(138, 43, 226, 0.5); background-color: rgba(0,0,0,0.2);
    color: var(--text-color); font-size: 1rem; resize: vertical;
    transition: border-color 0.3s, box-shadow 0.3s;
}
input:focus, textarea:focus { outline: none; border-color: var(--secondary-accent); box-shadow: 0 0 15px var(--glow-color); }
.autofill-section { display: flex; gap: 15px; align-items: center; margin-top: 15px;}
button {
    background: linear-gradient(90deg, var(--primary-accent), var(--secondary-accent));
    color: white; padding: 15px 30px; border: none; border-radius: 10px; cursor: pointer;
    font-weight: 600; font-size: 1rem; transition: all 0.3s ease;
}
button:hover { transform: scale(1.05); box-shadow: 0 5px 20px var(--glow-color); }
button.search-db-btn {
    padding: 2px 8px;
    font-size: 0.8rem;
    font-weight: 400;
    margin-left: 10px;
    border-radius: 5px;
    background: var(--secondary-accent);
}
.dashboard-header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 20px; padding: 0 10px; }
.dashboard-header h2 { color: #fff; }
#downloadReportBtn { background: linear-gradient(90deg, #3f9a76, #2a6f53); }
.dashboard-grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(500px, 1fr)); gap: 30px; }
.full-span-card { grid-column: 1 / -1; }
.card-header { margin-top: 0; border-bottom: 1px solid rgba(138, 43, 226, 0.5); padding-bottom: 15px; font-size: 1.4rem; color: #fff; }
.chat-box { height: 300px; overflow-y: auto; background-color: rgba(0,0,0,0.2); padding: 15px; border-radius: 10px; }
.qa-input { display: flex; gap: 10px; margin-top: 15px; }
.user-msg, .ai-msg { padding: 10px 15px; border-radius: 18px; margin-bottom: 10px; max-width: 85%; animation: popIn 0.3s ease-out; }
.user-msg { background: var(--primary-accent); align-self: flex-end; }
.ai-msg { background: #2c3e50; align-self: flex-start; }
.risk-item { padding: 15px; border-radius: 10px; margin: 15px 0; border-left: 5px solid; transition: transform 0.3s; }
.risk-high { border-color: #e94560; background-color: var(--high-risk-bg); }
.risk-caution { border-color: #f7b731; background-color: var(--caution-bg); }
.code-box, .result-box { margin-top: 15px; padding: 15px; background-color: rgba(0,0,0,0.2); border-radius: 10px; white-space: pre-wrap; }
.result-box ul { padding-left: 20px; }
footer { text-align: center; margin-top: 40px; color: var(--text-muted); font-size: 0.9rem; }
.hidden { display: none !important; }
#loader { border: 8px solid var(--primary-accent); border-top: 8px solid var(--secondary-accent); border-radius: 50%; width: 60px; height: 60px; animation: spin 1s linear infinite; margin: 40px auto; }
@keyframes spin { 100% { transform: rotate(360deg); } }
@keyframes pulse { 0%, 100% { transform: scale(1); } 50% { transform: scale(1.1); } }
@keyframes card-fade-in { from { opacity: 0; transform: translateY(20px); } to { opacity: 1; transform: translateY(0); } }
"""

JS_CODE = """
document.addEventListener('DOMContentLoaded', () => {
    // --- Particle Background ---
    const canvas = document.getElementById('particle-canvas');
    const ctx = canvas.getContext('2d');
    canvas.width = window.innerWidth;
    canvas.height = window.innerHeight;
    let particles = [];
    function initParticles() {
        particles = [];
        for (let i = 0; i < 50; i++) {
            particles.push({ x: Math.random() * canvas.width, y: Math.random() * canvas.height, size: Math.random() * 2 + 1, speedX: Math.random() * 1 - 0.5, speedY: Math.random() * 1 - 0.5, color: `rgba(138, 43, 226, ${Math.random()})` });
        }
    }
    function animateParticles() {
        ctx.clearRect(0, 0, canvas.width, canvas.height);
        for (const p of particles) {
            if (p.x > canvas.width || p.x < 0) p.speedX *= -1;
            if (p.y > canvas.height || p.y < 0) p.speedY *= -1;
            p.x += p.speedX; p.y += p.speedY;
            ctx.fillStyle = p.color;
            ctx.beginPath();
            ctx.arc(p.x, p.y, p.size, 0, Math.PI * 2);
            ctx.fill();
        }
        requestAnimationFrame(animateParticles);
    }
    initParticles();
    animateParticles();
    window.addEventListener('resize', () => { canvas.width = window.innerWidth; canvas.height = window.innerHeight; initParticles(); });

    // --- Splash Screen ---
    const splash = document.getElementById('splash-screen');
    setTimeout(() => {
        splash.style.opacity = '0';
        document.getElementById('app-body').classList.remove('hidden');
        setTimeout(() => splash.classList.add('hidden'), 500);
    }, 1500);

    // --- Global State & Elements ---
    let currentDocumentText = '', negotiationHistory = [];
    const langSelector = document.getElementById('languageSelector');

    // --- Helper Functions ---
    async function fetchAPI(endpoint, body, method = 'POST', stringify = true) {
        try {
            const options = { method };
            if (method !== 'GET') {
                if (stringify) {
                    options.headers = { 'Content-Type': 'application/json' };
                    options.body = JSON.stringify(body);
                } else {
                    options.body = body;
                }
            }
            const response = await fetch(endpoint, options);
            if (!response.ok) {
                const err = await response.json().catch(() => ({ detail: 'Server error' }));
                throw new Error(err.detail);
            }
            return await response.json();
        } catch (error) {
            console.error(`API Error (${endpoint}):`, error);
            alert(`An error occurred: ${error.message}`);
            return null;
        }
    }

    function typewriterEffect(element, text) { let i = 0; element.innerHTML = ''; const cursor = document.createElement('span'); cursor.className = 'cursor'; element.appendChild(cursor); function type() { if (i < text.length) { element.insertBefore(document.createTextNode(text.charAt(i)), cursor); i++; setTimeout(type, 20); } else { cursor.remove(); } } type(); }

    // --- Event Listeners ---
    document.getElementById('documentUpload').addEventListener('change', (e) => {
        const file = e.target.files[0];
        document.getElementById('fileName').textContent = file ? file.name : '';
    });

    document.getElementById('analyzeBtn').addEventListener('click', async () => {
        const file = document.getElementById('documentUpload').files[0];
        if (!file) { alert("Please upload a document."); return; }
        document.getElementById('loader').classList.remove('hidden');
        document.getElementById('dashboard').classList.add('hidden');
        const formData = new FormData();
        formData.append('file', file);
        formData.append('user_name', document.getElementById('autofillName').value);
        formData.append('language', langSelector.value);
        const data = await fetchAPI('/analyze', formData, 'POST', false);
        document.getElementById('loader').classList.add('hidden');
        if (data) {
            currentDocumentText = data.original_document;
            document.getElementById('keyFactsOutput').innerHTML = data.key_facts;
            document.getElementById('riskAnalysisOutput').innerHTML = data.risk_analysis;
            document.getElementById('lifespanOutput').innerHTML = data.lifespan;
            document.getElementById('filledDocumentOutput').textContent = data.filled_document;
            negotiationHistory = data.negotiation_history || [];
            document.getElementById('negotiationOutput').innerHTML = `<div class="ai-msg">${data.negotiation_start || ''}</div>`;
            document.getElementById('dashboard').classList.remove('hidden');
        }
    });

    const legalAidOutput = document.getElementById('legalAidOutput');
    async function performAndRenderLawSearch(payload) {
        legalAidOutput.innerHTML = '<em>Fetching relevant laws and suggestions...</em>';
        const response = await fetchAPI('/laws_search', payload, 'POST');

        if (!response) {
            legalAidOutput.innerHTML = '<p>Error fetching data.</p>';
            return;
        }

        let finalHtml = response.laws_html || "";
        if (response.ai_suggestions_json) {
            finalHtml += "<h4 style='margin-top:12px'>AI Suggestions</h4>";
            try {
                const cleanJsonString = response.ai_suggestions_json.replace(/```json\\n?|\\n?```/g, '').trim();
                const suggestionsData = JSON.parse(cleanJsonString);
                
                if (suggestionsData.laws && suggestionsData.laws.length > 0) {
                    finalHtml += "<h5>Relevant Legal Topics</h5><ul>";
                    suggestionsData.laws.forEach(law => {
                        finalHtml += `<li><strong>${law.title}:</strong> ${law.reason} <button class="search-db-btn" data-law-title="${law.title}">Search DB</button></li>`;
                    });
                    finalHtml += "</ul>";
                }
                if (suggestionsData.suggestions && suggestionsData.suggestions.length > 0) {
                    finalHtml += "<h5 style='margin-top:10px;'>Actionable Suggestions</h5><ul>";
                    suggestionsData.suggestions.forEach(suggestion => {
                        finalHtml += `<li>${suggestion}</li>`;
                    });
                    finalHtml += "</ul>";
                }
            } catch (e) {
                finalHtml += `<p>Could not format AI suggestions. Raw response:</p><pre>${response.ai_suggestions_json}</pre>`;
            }
        }
        legalAidOutput.innerHTML = finalHtml;
    }
    
    document.getElementById('legalAidBtn').addEventListener('click', () => {
        performAndRenderLawSearch({ document_text: currentDocumentText, q: "", language: langSelector.value, jurisdiction: "Chennai" });
    });

    legalAidOutput.addEventListener('click', (event) => {
        if (event.target.classList.contains('search-db-btn')) {
            const lawTitle = event.target.getAttribute('data-law-title');
            performAndRenderLawSearch({ document_text: "", q: lawTitle, language: langSelector.value, jurisdiction: "India" });
        }
    });
});
"""

# --- 6. STARTUP/SHUTDOWN & APP INITIALIZATION ---
@asynccontextmanager
async def lifespan(app: FastAPI):
    print("🚀 Running application startup logic...")
    db = SessionLocal()
    try:
        seed_sample_laws(db)
    finally:
        db.close()
    print("✅ Application startup complete.")
    yield
    print("🌙 Application shutting down.")

app = FastAPI(title="Legal AI Dashboard (Final)", lifespan=lifespan)

# --- 7. API ROUTES ---
@app.get("/", response_class=HTMLResponse)
async def home():
    return HTML_TEMPLATE.replace("{{ css_code }}", CSS_CODE).replace("{{ js_code }}", JS_CODE)

@app.post("/analyze")
async def analyze_document(user_name: str = Form(""), language: str = Form(...), file: UploadFile = File(...)):
    doc_text = extract_text_from_file(file)
    filled_doc = doc_text.replace("[Your Name]", user_name or "[[Your Name]]")
    lang_name = language or "English"
    risk_prompt = f"""Analyze all potential risks in the legal document. For each risk, create an HTML element.
- For HIGH risks (major financial loss, liability), use: <div class='risk-item risk-high'><strong>High Risk:</strong> [Description]</div>
- For MODERATE risks (unfavorable terms, negotiation points), use: <div class='risk-item risk-caution'><strong>Moderate Risk:</strong> [Description]</div>
Respond only with HTML in {lang_name}. Document:\n{filled_doc}"""
    prompts = [
        f"Extract key facts (parties, dates, amounts) as HTML. Respond in {lang_name}.\n\n{filled_doc}",
        risk_prompt,
        f"Create a legal timeline (key dates, deadlines) as HTML. Respond in {lang_name}.\n\n{filled_doc}"
    ]
    key, risk, life = await asyncio.gather(*[get_ai_response(p) for p in prompts])
    neg_hist = [{'role': 'user', 'parts': [f"You are a Landlord in Chennai. I am a Tenant. Start negotiating this document in {lang_name}.\n\n{filled_doc}"]}]
    init_response_text = "AI negotiation is unavailable."
    try:
        if API_KEY:
            chat = model.start_chat(history=neg_hist)
            init_response = await chat.send_message_async("Start now.")
            init_response_text = init_response.text
            neg_hist.append({'role': 'model', 'parts': [init_response_text]})
    except Exception as e:
        init_response_text = f"AI negotiation failed: {e}"
    return {'key_facts': key, 'risk_analysis': risk, 'lifespan': life, 'original_document': doc_text, 'filled_document': filled_doc, 'negotiation_start': init_response_text, 'negotiation_history': neg_hist}

@app.post("/laws_search")
async def laws_search(req: LawSearchRequest, db: Session = Depends(get_db)):
    q = (req.q or "").strip()
    jurisdiction = (req.jurisdiction or "").strip()
    if not q and req.document_text:
        words = re.findall(r"\b[a-zA-Z]{5,}\b", req.document_text.lower())
        top_words = sorted({w: words.count(w) for w in set(words)}.items(), key=lambda item: item[1], reverse=True)[:5]
        q = " ".join([word for word, count in top_words])

    laws = simple_law_search(db, q=q or None, jurisdiction=jurisdiction or None, limit=10)
    laws_html = "".join([f"<div class='risk-item'><strong>{law.title}</strong><br/><small>{law.jurisdiction} — tags: {law.tags}</small><div>{law.text[:800]}{'...' if len(law.text) > 800 else ''}</div></div>" for law in laws])
    
    ai_raw_json = None
    if req.document_text and API_KEY and not req.q:
        prompt = (f"You are a legal assistant. For the document excerpt below, identify relevant legal topics, explain why, "
                  f"and propose 2 practical suggestions for someone in Chennai. Respond in {req.language or 'English'}.\n\n"
                  f"Document excerpt:\n{req.document_text[:2000]}\n\n"
                  f"Respond ONLY in this JSON format: {{\"laws\": [{{\"title\": \"Law Title\", \"reason\": \"Explanation\"}}], \"suggestions\": [\"Suggestion 1\", \"Suggestion 2\"]}}")
        ai_raw_json = await get_ai_response(prompt)
        try:
            db.add(SuggestionRecord(source_doc_preview=req.document_text[:400], suggested_text=ai_raw_json[:2000], reason="AI law-match"))
            db.commit()
        except Exception as e:
            print("Suggestion store error:", e)

    return JSONResponse({
        "laws_html": "<h4>Relevant Laws from Database</h4>" + (laws_html or "<p>No matching laws found.</p>"),
        "ai_suggestions_json": ai_raw_json
    })

# (Other endpoints like compare_clauses, download_report, etc.)
@app.post("/compare_clauses")
async def compare_clauses(req: CompareRequest):
    prompt = f"Compare Clause A and B. Respond in {req.language} with an HTML list of differences and risks.\nA:{req.clause_a}\nB:{req.clause_b}"
    return {'comparison': await get_ai_response(prompt)}

@app.post("/download_report")
async def download_report(req: ReportRequest):
    pdf = PDF()
    pdf.add_page()
    pdf.chapter_title("1. Key Facts & Figures")
    pdf.write_html_content(req.key_facts)
    pdf.chapter_title("2. Risk Analysis")
    pdf.write_html_content(req.risk_analysis)
    pdf.chapter_title("3. Auto-Filled Document Text")
    pdf.set_font("Courier", "", 9)
    pdf.multi_cell(0, 5, sanitize_for_pdf(req.filled_document))
    pdf_bytes = pdf.output(dest="S").encode("latin-1", "replace")
    return StreamingResponse(io.BytesIO(pdf_bytes), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=Legal_AI_Report.pdf"})

@app.post("/ask")
async def ask_question(req: AskRequest):
    prompt = f"Answer in {req.language}. DOC:\n{req.document_text}\n\nQ:\n{req.question}"
    return {'answer': await get_ai_response(prompt)}

@app.post("/negotiate")
async def negotiate(req: NegotiateRequest):
    history = req.history or []
    history.append({'role': 'user', 'parts': [req.user_message]})
    try:
        chat = model.start_chat(history=history)
        resp = await chat.send_message_async("Continue negotiation.")
        history.append({'role': 'model', 'parts': [resp.text]})
        return {'ai_response': resp.text, 'updated_history': history}
    except Exception as e:
        return {'ai_response': f"AI error: {e}", 'updated_history': history}

# --- 8. RUN APP ---
if __name__ == "__main__":
   
    uvicorn.run(app, host="0.0.0.0", port=8000)